import sys
sys.path.insert(0, "/tmp/pptx_lib")
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION

FONT = "Leelawadee UI"
NAVY = RGBColor(0x1F, 0x2A, 0x44)
ORANGE = RGBColor(0xE8, 0x6A, 0x1F)
TEAL = RGBColor(0x1B, 0x8A, 0x8F)
LIGHT = RGBColor(0xF4, 0xF6, 0xFA)
GRAY = RGBColor(0x5A, 0x63, 0x73)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xC0, 0x39, 0x2B)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, s, size=18, bold=False, color=NAVY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = FONT
        p.space_after = Pt(6)
    return tb


def box(slide, x, y, w, h, fill=LIGHT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1.5)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def card(slide, x, y, w, h, title, body, accent=TEAL, size=15):
    box(slide, x, y, w, h, LIGHT)
    box(slide, x, y, 0.09, h, accent, shape=MSO_SHAPE.RECTANGLE)
    text(slide, x + 0.25, y + 0.1, w - 0.4, 0.5, title, size=size + 3, bold=True, color=accent)
    text(slide, x + 0.25, y + 0.62, w - 0.4, h - 0.7, body, size=size, color=GRAY)


def header(slide, title, sub=None):
    box(slide, 0, 0, 13.333, 1.15, NAVY, shape=MSO_SHAPE.RECTANGLE)
    box(slide, 0, 1.15, 13.333, 0.06, ORANGE, shape=MSO_SHAPE.RECTANGLE)
    text(slide, 0.6, 0.18, 12, 0.7, title, size=30, bold=True, color=WHITE)
    if sub:
        text(slide, 0.6, 0.72, 12, 0.4, sub, size=14, color=RGBColor(0xC9, 0xD1, 0xE3))


def footer(slide, n):
    text(slide, 0.6, 7.05, 9, 0.3, "OCR & LLM Curriculum QA  |  06026240 Intelligent System Development", size=10, color=GRAY)
    text(slide, 12.0, 7.05, 0.8, 0.3, str(n), size=10, color=GRAY, align=PP_ALIGN.RIGHT)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def table(slide, x, y, w, h, rows, col_w=None, size=14):
    t = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci)
            c.text = ""
            r = c.text_frame.paragraphs[0].add_run()
            r.text = str(val)
            r.font.size = Pt(size)
            r.font.name = FONT
            r.font.bold = ri == 0
            r.font.color.rgb = WHITE if ri == 0 else NAVY
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if ri == 0 else (LIGHT if ri % 2 else WHITE)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
    if col_w:
        for i, cw in enumerate(col_w):
            t.columns[i].width = Inches(cw)
    return t


def new(title, sub=None):
    s = prs.slides.add_slide(BLANK)
    header(s, title, sub)
    footer(s, len(prs.slides))
    return s

# ---------- 1. Title ----------
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, NAVY, shape=MSO_SHAPE.RECTANGLE)
box(s, 0.8, 3.55, 2.2, 0.08, ORANGE, shape=MSO_SHAPE.RECTANGLE)
text(s, 0.8, 1.7, 11.5, 1.8, ["ระบบถาม-ตอบเล่มหลักสูตร", "ด้วย OCR + Local LLM"], size=44, bold=True, color=WHITE)
text(s, 0.8, 3.8, 11, 0.6, "ถามเป็นภาษาไทย ตอบจากฐานข้อมูลหลักสูตรจริง ไม่เดา ไม่แต่งเรื่อง", size=22, color=RGBColor(0xC9, 0xD1, 0xE3))
text(s, 0.8, 6.2, 11, 0.8, ["06026240 Intelligent System Development", "ภาควิชาเทคโนโลยีสารสนเทศ คณะเทคโนโลยีสารสนเทศ สจล."], size=16, color=RGBColor(0xC9, 0xD1, 0xE3))
notes(s, "เปิดด้วยปัญหา: นักศึกษาต้องเปิด PDF หลักสูตรหลายร้อยหน้าเพื่อหาว่าวิชานี้กี่หน่วยกิต เรียนเทอมไหน ต้องเรียนอะไรก่อน")

# ---------- 2. Problem ----------
s = new("ปัญหา: เล่มหลักสูตรหาคำตอบยาก", "มคอ.2 เป็น PDF หลายร้อยหน้า ตารางซับซ้อน ฟอนต์ไทยเพี้ยน")
card(s, 0.6, 1.6, 3.9, 2.6, "หาข้อมูลยาก", ["\"วิชานี้เรียนเทอมไหน\"", "\"ต้องเรียนอะไรมาก่อน\"", "ต้องไล่เปิดหลายหน้า"], ORANGE, 16)
card(s, 4.7, 1.6, 3.9, 2.6, "ให้ LLM อ่านตรง ๆ ก็ไม่พอ", ["นับหน่วยกิตผิดบ่อย", "ตกหล่นเมื่อข้อมูลข้ามหน้า", "ตอบผิดแบบดูน่าเชื่อถือ"], RED, 16)
card(s, 8.8, 1.6, 3.9, 2.6, "สิ่งที่เราต้องการ", ["คำตอบถูกต้องทุกครั้ง", "รู้ว่าเมื่อไหร่ไม่รู้", "รันในเครื่อง ไม่ส่งข้อมูลออก"], TEAL, 16)
box(s, 0.6, 4.7, 12.1, 1.6, NAVY)
text(s, 0.9, 4.85, 11.5, 1.3, ["แนวคิดหลัก", "ให้ LLM ทำสิ่งที่เก่ง (แปลภาษาคนเป็นคำสั่ง) ให้ฐานข้อมูลทำสิ่งที่เก่ง (นับ รวม เปรียบเทียบ)"], size=20, bold=True, color=WHITE)
notes(s, "LLM ไม่ใช่เครื่องคิดเลข เวลามันผิดมักผิดแบบมั่นใจมาก เราจึงย้ายงานคำนวณไปให้ SQL")

# ---------- 3. Pipeline ----------
s = new("ภาพรวมระบบ 4 ขั้นตอน", "จากไฟล์ PDF ถึงคำตอบภาษาไทย")
steps = [("1", "อ่านเอกสาร", "PDF → ข้อความ\nแก้ฟอนต์ MacThai\nOCR (Typhoon)"),
         ("2", "สกัดข้อมูล", "Qwen3:4b\nข้อความ → JSON\nตาม Schema"),
         ("3", "เก็บลงฐานข้อมูล", "SQLite 4 ตาราง\nตรวจความสอดคล้อง\n7 กฎ"),
         ("4", "ถาม-ตอบ", "ภาษาไทย → SQL\nรันบน DB\nสรุปเป็นภาษาไทย")]
for i, (n, t, b) in enumerate(steps):
    x = 0.6 + i * 3.15
    box(s, x, 1.9, 2.8, 3.4, LIGHT, line=TEAL)
    c = box(s, x + 0.95, 2.05, 0.9, 0.9, ORANGE, shape=MSO_SHAPE.OVAL)
    text(s, x + 0.95, 2.15, 0.9, 0.7, n, size=28, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, x + 0.1, 3.1, 2.6, 0.5, t, size=20, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    text(s, x + 0.1, 3.7, 2.6, 1.5, b.split("\n"), size=14, color=GRAY, align=PP_ALIGN.CENTER)
    if i < 3:
        text(s, x + 2.78, 3.2, 0.4, 0.5, "→", size=28, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)
text(s, 0.6, 5.7, 12, 0.9, "วัดผลแยกทีละขั้น เมื่อพังจะรู้ทันทีว่าพังตรงไหน (หลักจากบทที่ 9: Evaluation)", size=18, color=NAVY, align=PP_ALIGN.CENTER)
notes(s, "เน้นว่าแต่ละขั้นตอนมี metric ของตัวเอง OCR ใช้ CER/WER, Extraction ใช้ Recall/Precision, QA ใช้ Execution/Answer accuracy")

# SLIDES_PLACEHOLDER
